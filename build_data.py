"""Build the fixed MOH snapshot using openpyxl and Python's standard library.

Run: python build_data.py
Optional independent check: python build_data.py --verify-reference PATH

No network access, imputation, rounding, or patient-level inference is performed.
Source-cell references distinguish a blank cell from an absent source date.
"""

from __future__ import annotations

import argparse
from collections import Counter
from contextlib import closing
import csv
from datetime import date, datetime, timedelta
from hashlib import sha256
import json
import math
from pathlib import Path
import sqlite3
import sys
from tempfile import TemporaryDirectory

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


HERE = Path(__file__).resolve().parent
STUDY_START = date(2024, 1, 1)
STUDY_END = date(2026, 9, 12)
DOWNLOAD_DATE = "2026-09-20"
HOSPITALS = ["AH", "CGH", "KTPH", "NTFGH", "NUH(A)", "SGH", "SKH", "TTSH", "WH"]
FIELDS = [
    "date", "hospital_code", "ed_attendances", "bor_ratio",
    "admission_wait_p50_hours", "availability", "ed_source", "bor_source", "wait_source",
]
SPECS = [
    {
        "prefix": "ed", "metric": "ed_attendances",
        "file": "急诊就诊量_EMD_week36Y2026.xlsx", "history": "Historical",
        "sha256": "048ad30421329d43c39ae827afdeaa3b94df3af8c105d509db8f1ade881740d5",
        "url": "https://isomer-user-content.by.gov.sg/3/989bf3f6-f3ac-4e47-8ce7-9607b628a1af/Attendances%20at%20EMD_week36Y2026.xlsx",
    },
    {
        "prefix": "bor", "metric": "bor_ratio",
        "file": "床位占用率_BOR_week36Y2026.xlsx", "history": "BOR(%)_historical",
        "sha256": "51b7957349e3a86cfc2938deba1c33b88ce798109d6d53d23f14db52b289c882",
        "url": "https://isomer-user-content.by.gov.sg/3/64172059-8696-4555-b7de-e91b0a6110cb/Bed%20Occupancy%20Rate_week36Y2026.xlsx",
    },
    {
        "prefix": "wait", "metric": "admission_wait_p50_hours",
        "file": "入院等待时间_WT_week36Y2026.xlsx", "history": "Historical",
        "sha256": "a3a3e7aad1e78ea7bc9d368f25118b00ad36217a00bb90a200eb63d6bc8f5508",
        "url": "https://isomer-user-content.by.gov.sg/3/a919d85a-4350-4cd4-a4ff-cac891237d9d/WT%20for%20Admission%20to%20Ward_week36Y2026.xlsx",
    },
]


def study_dates():
    return [STUDY_START + timedelta(days=i) for i in range((STUDY_END - STUDY_START).days + 1)]


def require(condition, message):
    """Unlike assert, validation must also run under python -O."""
    if not condition:
        raise ValueError(message)


def read_metric(raw_dir, spec):
    path = raw_dir / spec["file"]
    digest = sha256(path.read_bytes()).hexdigest()
    require(digest == spec["sha256"], f"SHA-256 mismatch: {path}; review the source version before rebuilding")
    records = {}
    sheet_counts = []
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in (spec["history"], "Sheet1"):
            require(sheet_name in book.sheetnames, f"Missing sheet: {path.name} / {sheet_name}")
            sheet = book[sheet_name]
            header = next(sheet.iter_rows(min_row=3, max_row=3, max_col=10, values_only=True))
            require(list(header) == ["Date", *HOSPITALS], f"Unexpected header: {path.name} / {sheet_name}")
            date_count = 0
            for source_row, cells in enumerate(sheet.iter_rows(min_row=4, max_col=10), 4):
                raw_day = cells[0].value
                if not isinstance(raw_day, (date, datetime)):
                    continue  # Titles and explanatory notes are not observation rows.
                day = raw_day.date() if isinstance(raw_day, datetime) else raw_day
                if not STUDY_START <= day <= STUDY_END:
                    continue
                date_count += 1
                for source_column, (hospital, cell) in enumerate(zip(HOSPITALS, cells[1:]), 2):
                    key = (day.isoformat(), hospital)
                    # Read-only EmptyCell has no coordinate; row/column still identify the original blank.
                    source = f"{sheet_name}!{get_column_letter(source_column)}{source_row}"
                    require(key not in records, f"Duplicate date/hospital: {path.name} / {key}")
                    value = cell.value
                    if value is not None:
                        valid = not isinstance(value, bool) and isinstance(value, (int, float)) and math.isfinite(value)
                        require(valid, f"Non-finite or non-numeric value: {path.name} / {source}")
                        require(value >= 0, f"Negative value requires source review: {path.name} / {source}")
                        if spec["prefix"] == "ed":
                            require(value == int(value), f"ED count is not an integer: {path.name} / {source}")
                            value = int(value)
                        if spec["prefix"] == "bor":
                            require("%" in cell.number_format, f"BOR percentage format changed: {path.name} / {source}")
                            # Keep 0.95 as 0.95. Do not clip possible values above 1.
                    records[key] = (value, source)
            sheet_counts.append({"sheet": sheet_name, "studyDateRows": date_count})
    finally:
        book.close()
    require(sha256(path.read_bytes()).hexdigest() == digest, f"Source changed during reading: {path}")
    dates_present = {key[0] for key in records}
    values = [value for value, _ in records.values() if value is not None]
    audit = {
        "metric": spec["metric"], "file": path.name, "sha256": digest,
        "sheetCounts": sheet_counts, "dateCount": len(dates_present), "rowCount": len(records),
        "nonblankCount": len(values), "blankCount": len(records) - len(values),
        "zeroCount": sum(value == 0 for value in values),
        "missingDates": [day.isoformat() for day in study_dates() if day.isoformat() not in dates_present],
        "blankByHospital": {h: sum(v is None for (d, hospital), (v, s) in records.items() if hospital == h) for h in HOSPITALS},
    }
    if spec["prefix"] == "bor":
        audit["above100PercentCount"] = sum(value > 1 for value in values)
    return records, audit


def availability(values):
    labels = [label for value, label in zip(values, ["急诊量", "床位占用率", "等待时间"]) if value is not None]
    if len(labels) == 3:
        return "三项有值"
    return "仅" + "、".join(labels) + "有值" if labels else "三项无数值"


def build_payload(raw_dir):
    metrics, audits = {}, []
    for spec in SPECS:
        metrics[spec["prefix"]], audit = read_metric(raw_dir, spec)
        audits.append(audit)
    # Full outer join: use the union of source keys, including original blank cells.
    union_keys = set().union(*(set(records) for records in metrics.values()))
    rows = []
    for day, hospital in sorted(union_keys):
        source_records = [metrics[prefix].get((day, hospital), (None, None)) for prefix in ("ed", "bor", "wait")]
        values = [record[0] for record in source_records]
        sources = [record[1] for record in source_records]
        rows.append([day, hospital, *values, availability(values), *sources])
    keys = [(row[0], row[1]) for row in rows]
    require(len(keys) == len(set(keys)), "Joined keys are duplicated")
    expected_keys = {(day.isoformat(), h) for day in study_dates() for h in HOSPITALS}
    require(set(keys) == expected_keys, "Snapshot does not cover the expected date/hospital grid")
    by_key = {tuple(row[:2]): row for row in rows}
    for metric_index, prefix in enumerate(("ed", "bor", "wait")):
        for key, (value, source) in metrics[prefix].items():
            row = by_key[key]
            require((row[2 + metric_index], row[6 + metric_index]) == (value, source), f"Join changed value or source: {prefix} / {key}")
    counts = Counter(row[5] for row in rows)
    quality = {
        "dateCount": len({row[0] for row in rows}), "rowCount": len(rows),
        "completeCount": counts["三项有值"], "borOnlyCount": counts["仅床位占用率有值"],
        "allMissingCount": counts["三项无数值"], "duplicateKeys": len(keys) - len(set(keys)),
    }
    expected_quality = {"dateCount": 986, "rowCount": 8874, "completeCount": 8637, "borOnlyCount": 63, "allMissingCount": 174, "duplicateKeys": 0}
    require(quality == expected_quality, f"Unexpected snapshot quality: {quality}")
    payload = {
        "fields": FIELDS, "rows": rows,
        "meta": {
            "studyStart": STUDY_START.isoformat(), "studyEnd": STUDY_END.isoformat(),
            "downloadDate": DOWNLOAD_DATE, "hospitals": HOSPITALS,
            "sources": [{key: spec[key] for key in ("metric", "file", "url", "sha256")} for spec in SPECS],
        },
        "quality": quality,
    }
    audit = {
        "quality": quality, "sources": audits,
        "availabilityCounts": dict(counts),
        "join": "one-to-one full outer join on date + unchanged hospital_code",
        "missingPolicy": "Null stays null; an absent source date has a null source; a blank source cell keeps its source coordinate. Zero stays zero.",
        "borStorage": "Unrounded original ratio; display 0.95 as 95%, without altering storage.",
        "limitations": "Same-day alignment does not identify the same patients or a causal relationship. No patient-level data or exact daily bed denominator is available.",
    }
    return payload, audit


def verify_reference(payload, reference_path):
    reference = json.loads(reference_path.read_text(encoding="utf-8"))
    joined = reference.get("joined", reference)
    require(payload["fields"] == joined["fields"], "Reference field order differs")
    require(len(payload["rows"]) == len(joined["rows"]), "Reference row count differs")
    for index, (actual, expected) in enumerate(zip(payload["rows"], joined["rows"]), 1):
        require(actual == expected, f"Reference row {index} differs: {actual[:2]}")
    return {"matchedRows": len(payload["rows"]), "matchedFieldsPerRow": len(FIELDS), "valuesAndSourcesEqual": True}


def write_sqlite(path, rows):
    with closing(sqlite3.connect(path)) as connection:
        connection.execute("""
            CREATE TABLE hospital_metrics (
                date TEXT NOT NULL,
                hospital_code TEXT NOT NULL,
                ed_attendances INTEGER,
                bor_ratio REAL,
                admission_wait_p50_hours REAL,
                availability TEXT NOT NULL,
                ed_source TEXT,
                bor_source TEXT,
                wait_source TEXT,
                PRIMARY KEY (date, hospital_code),
                CHECK (ed_attendances IS NULL OR (typeof(ed_attendances) = 'integer' AND ed_attendances >= 0)),
                CHECK (bor_ratio IS NULL OR bor_ratio >= 0),
                CHECK (admission_wait_p50_hours IS NULL OR admission_wait_p50_hours >= 0)
            )
        """)
        connection.executemany("INSERT INTO hospital_metrics VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)", rows)
        connection.commit()
        require(connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok", "SQLite integrity check failed")
        restored = connection.execute("SELECT * FROM hospital_metrics ORDER BY date, hospital_code").fetchall()
        require([list(row) for row in restored] == rows, "SQLite round-trip changed a value or source")


def export(payload, audit, output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    # Prepare every artifact in the output directory before replacing existing files.
    with TemporaryDirectory(prefix=".build-", dir=output_dir) as work:
        stage = Path(work)
        text = json.dumps(payload, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
        (stage / "dashboard-data.json").write_text(text + "\n", encoding="utf-8")
        (stage / "dashboard-data.js").write_text("window.HOSPITAL_DATA=" + text + ";\n", encoding="utf-8")
        with (stage / "hospital_metrics.csv").open("w", encoding="utf-8-sig", newline="") as file:
            writer = csv.writer(file)
            writer.writerow(FIELDS)
            writer.writerows(payload["rows"])
        write_sqlite(stage / "hospital_metrics.sqlite", payload["rows"])
        (stage / "build-audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
        for path in sorted(stage.iterdir()):
            path.replace(output_dir / path.name)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=HERE / "raw_data")
    parser.add_argument("--output-dir", type=Path, default=HERE / "data")
    parser.add_argument("--verify-reference", type=Path, help="Optional legacy JSON for exact value/source comparison; not a build input")
    args = parser.parse_args()
    payload, audit = build_payload(args.raw_dir)
    reference_result = verify_reference(payload, args.verify_reference) if args.verify_reference else None
    export(payload, audit, args.output_dir)
    result = {"quality": payload["quality"], "outputDirectory": str(args.output_dir.resolve())}
    if reference_result is not None:
        result["referenceVerification"] = reference_result
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
