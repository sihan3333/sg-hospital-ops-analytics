"""独立从原始 XLSX 逐格检查已导出的长表和三指标合并表。"""
from collections import Counter
from datetime import datetime
from hashlib import sha256
import json
import math
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / '原始数据/2026-09-20'
WORK = ROOT / '处理脚本/.work_join'
OUTPUT = ROOT / 'outputs/01a0bcf0-303c-7d21-9f82-aafceef76c21/医疗运营_三指标合并.xlsx'
metadata = json.loads((WORK / 'hospital_metrics.json').read_text(encoding='utf-8'))
sources = {}
for prefix in ('ed', 'bor', 'wait'):
    spec = metadata['datasets'][prefix]['audit']
    source_path = RAW / spec['file']
    assert sha256(source_path.read_bytes()).hexdigest() == spec['sha256']
    book = load_workbook(source_path, read_only=False, data_only=True)
    original = {}
    for sheet in book:
        hospitals = [sheet.cell(3, col).value for col in range(2, 11)]
        for row in range(4, sheet.max_row + 1):
            day = sheet.cell(row, 1).value
            if not isinstance(day, datetime) or day < datetime(2024, 1, 1):
                continue
            for col, hospital in enumerate(hospitals, 2):
                cell = sheet.cell(row, col)
                key = (day.date().isoformat(), hospital)
                assert key not in original
                original[key] = (cell.value, sheet.title, cell.coordinate)
    sources[prefix] = original
    book.close()


def same_number(actual, expected):
    if expected is None:
        assert actual is None
    else:
        assert actual is not None and math.isclose(actual, expected, rel_tol=1e-14, abs_tol=1e-14), (actual, expected)


book = load_workbook(OUTPUT, read_only=True, data_only=True)
long_checks = {}
for prefix, name in [('ed', '急诊量长表'), ('bor', '床位占用率长表')]:
    sheet = book[name]
    seen = set()
    for day, hospital, value, status, source_sheet, source_cell in sheet.iter_rows(min_row=12, max_col=6, values_only=True):
        key = (day.date().isoformat(), hospital)
        assert key not in seen
        seen.add(key)
        expected, expected_sheet, expected_cell = sources[prefix][key]
        same_number(value, expected)
        assert (source_sheet, source_cell) == (expected_sheet, expected_cell)
        assert status == ('原表空白' if expected is None else '原表有值')
    assert seen == set(sources[prefix])
    cached = [sheet[cell].value for cell in ['B5', 'D5', 'F5']]
    assert cached == ([8811, 8637, 174] if prefix == 'ed' else [8874, 8700, 174])
    long_checks[prefix] = {'keys_and_values_match': len(seen), 'cached_counts': cached}

sheet = book['三指标合并']
seen = set()
status_counts = Counter()
field_missing = Counter()
for row in sheet.iter_rows(min_row=12, max_col=9, values_only=True):
    day, hospital, ed, bor, wait, status, ed_location, bor_location, wait_location = row
    key = (day.date().isoformat(), hospital)
    assert key not in seen
    seen.add(key)
    for prefix, value, location in [('ed', ed, ed_location), ('bor', bor, bor_location), ('wait', wait, wait_location)]:
        if key in sources[prefix]:
            expected, expected_sheet, expected_cell = sources[prefix][key]
            same_number(value, expected)
            assert location == expected_sheet + '!' + expected_cell
        else:
            assert value is None and location is None
        if value is None:
            field_missing[prefix] += 1
    if ed is not None and bor is not None and wait is not None:
        assert status == '三项有值'
    elif ed is None and bor is None and wait is None:
        assert status == '三项无数值'
    else:
        assert ed is None and bor is not None and wait is None and status == '仅床位占用率有值'
    status_counts[status] += 1

assert seen == set().union(*(set(rows) for rows in sources.values()))
assert len(seen) == 8874
assert field_missing == {'ed': 237, 'bor': 174, 'wait': 237}
assert status_counts == {'三项有值': 8637, '三项无数值': 174, '仅床位占用率有值': 63}
assert [sheet[c].value for c in ['B5', 'D5', 'F5']] == [8874, 8637, 237]
assert sheet['D13'].number_format == '0.0%'
assert math.isclose(sheet['D13'].value, 0.8204121687929342)
book.close()
report = {'output': str(OUTPUT), 'long_tables': long_checks,
          'main_rows_verified_against_original': len(seen), 'status_counts': dict(status_counts),
          'missing_values_by_metric': dict(field_missing), 'original_hashes_unchanged': True,
          'bor_stored_ratio_and_percentage_format_verified': True}
(WORK / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=True, indent=2))
