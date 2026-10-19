"""三指标按官方所属日期和医院编码一对一外连接；保留缺行、空格和来源。"""
from collections import Counter
from datetime import date, datetime, timedelta
from hashlib import sha256
import json
import math
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / '原始数据/2026-09-20'
WORK = ROOT / '处理脚本/.work_join'
START = date(2024, 1, 1)
HOSPITALS = ['AH', 'CGH', 'KTPH', 'NTFGH', 'NUH(A)', 'SGH', 'SKH', 'TTSH', 'WH']
SPECS = [
    ('ed', 'ed_attendances', '急诊就诊量_EMD_week36Y2026.xlsx', 'Historical'),
    ('bor', 'bor_ratio', '床位占用率_BOR_week36Y2026.xlsx', 'BOR(%)_historical'),
    ('wait', 'admission_wait_p50_hours', '入院等待时间_WT_week36Y2026.xlsx', 'Historical'),
]


def json_rows(frame, fields):
    return [[None if pd.isna(v) else v for v in r] for r in frame[fields].itertuples(index=False, name=None)]


def read_metric(prefix, metric, filename, history_sheet):
    path = RAW / filename
    digest = sha256(path.read_bytes()).hexdigest()
    book = load_workbook(path, read_only=True, data_only=True)
    pieces, sheet_counts = [], []
    for sheet_name in (history_sheet, 'Sheet1'):
        sheet = book[sheet_name]
        header = [cell.value for cell in next(sheet.iter_rows(min_row=3, max_row=3, max_col=10))]
        if header != ['Date', *HOSPITALS]:
            raise ValueError(f'{filename} / {sheet_name}: 医院表头变化，停止自动匹配')
        wide_rows = []
        for source_row, cells in enumerate(sheet.iter_rows(min_row=4, max_col=10), 4):
            raw_date = cells[0].value
            if not isinstance(raw_date, (date, datetime)):
                continue
            day = raw_date.date() if isinstance(raw_date, datetime) else raw_date
            if day < START:
                continue
            values = [cell.value for cell in cells[1:]]
            for cell, value in zip(cells[1:], values):
                if value is None:
                    continue
                if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                    raise ValueError(f'{filename} / {sheet_name}!{cell.coordinate}: 非有限数值')
                if prefix == 'bor' and '%' not in cell.number_format:
                    raise ValueError(f'{sheet_name}!{cell.coordinate}: BOR 百分比格式变化，需核查单位')
                if prefix == 'ed' and value != int(value):
                    raise ValueError(f'{sheet_name}!{cell.coordinate}: 就诊计数不是整数，需核查')
            wide_rows.append([day.isoformat(), *values, sheet_name, source_row])
        sheet_counts.append({'sheet': sheet_name, 'study_date_rows': len(wide_rows)})
        frame = pd.DataFrame(wide_rows, columns=['date', *HOSPITALS, 'source_sheet', 'source_row'])
        pieces.append(frame.melt(id_vars=['date', 'source_sheet', 'source_row'], value_vars=HOSPITALS,
                                 var_name='hospital_code', value_name=metric))
    book.close()
    long = pd.concat(pieces, ignore_index=True).sort_values(['date', 'hospital_code'], kind='stable').reset_index(drop=True)
    if long.duplicated(['date', 'hospital_code']).any():
        raise ValueError(f'{filename}: 日期医院组合重复，必须先核查')
    columns = {hospital: chr(ord('B') + i) for i, hospital in enumerate(HOSPITALS)}
    long['source_cell'] = long.apply(lambda r: columns[r['hospital_code']] + str(int(r['source_row'])), axis=1)
    long['value_status'] = long[metric].apply(lambda v: '原表空白' if pd.isna(v) else '原表有值')
    long[prefix + '_source'] = long['source_sheet'] + '!' + long['source_cell']
    dates = set(long['date'])
    last = date.fromisoformat(max(dates))
    expected_dates = [(START + timedelta(days=i)).isoformat() for i in range((last - START).days + 1)]
    audit = {'file': filename, 'sha256': digest, 'metric': metric, 'sheet_counts': sheet_counts,
             'date_rows': len(dates), 'long_rows': len(long), 'nonblank_values': int(long[metric].notna().sum()),
             'blank_values': int(long[metric].isna().sum()), 'zero_values': int(long[metric].eq(0).sum()),
             'negative_values': int(long[metric].lt(0).sum()),
             'missing_whole_dates': [d for d in expected_dates if d not in dates],
             'blank_by_hospital': {h: int(long.loc[long['hospital_code'].eq(h), metric].isna().sum()) for h in HOSPITALS}}
    if prefix == 'bor':
        audit['above_100_percent_values'] = int(long[metric].gt(1).sum())
        audit['storage'] = '保持原始比例值，例如 0.95 在 Excel 中按 95.0% 显示；不乘100后再套百分比格式。'
    if sha256(path.read_bytes()).hexdigest() != digest:
        raise RuntimeError('读取期间原始文件版本发生改变')
    fields = ['date', 'hospital_code', metric, 'value_status', 'source_sheet', 'source_cell']
    artifact_data = {'fields': fields, 'rows': json_rows(long, fields), 'audit': audit}
    return long[['date', 'hospital_code', metric, prefix + '_source']], artifact_data


datasets, audits, frames = {}, {}, {}
for spec in SPECS:
    prefix = spec[0]
    frames[prefix], datasets[prefix] = read_metric(*spec)
    audits[prefix] = datasets[prefix]['audit']

keys = ['date', 'hospital_code']
# outer 保留任一来源存在的键；validate 防止意外一对多让行数膨胀。
joined = frames['wait'].merge(frames['ed'], on=keys, how='outer', validate='one_to_one')
joined = joined.merge(frames['bor'], on=keys, how='outer', validate='one_to_one')
joined = joined.sort_values(keys, kind='stable').reset_index(drop=True)
metrics = ['ed_attendances', 'bor_ratio', 'admission_wait_p50_hours']
names = dict(zip(metrics, ['急诊量', '床位占用率', '等待时间']))


def availability(row):
    available = [names[m] for m in metrics if pd.notna(row[m])]
    if len(available) == 3:
        return '三项有值'
    if not available:
        return '三项无数值'
    return '仅' + '、'.join(available) + '有值'


joined['availability'] = joined.apply(availability, axis=1)
expected_keys = set().union(*(set(zip(f['date'], f['hospital_code'])) for f in frames.values()))
assert set(zip(joined['date'], joined['hospital_code'])) == expected_keys
assert not joined.duplicated(keys).any()
# 按原长表的键回查，验证连接没有改变任何来源的数值、空白或出处。
for prefix, metric, _, _ in SPECS:
    source = frames[prefix].set_index(keys).sort_index()
    projected = joined.set_index(keys).loc[source.index, [metric, prefix + '_source']]
    pd.testing.assert_frame_equal(source[[metric, prefix + '_source']], projected, check_dtype=False)

join_fields = ['date', 'hospital_code', *metrics, 'availability', 'ed_source', 'bor_source', 'wait_source']
join_audit = {'date_rows': int(joined['date'].nunique()), 'hospital_codes': HOSPITALS,
              'study_start': joined['date'].min(), 'study_end': joined['date'].max(),
              'joined_rows': len(joined), 'complete_value_rows': int(joined[metrics].notna().all(axis=1).sum()),
              'availability_counts': dict(Counter(joined['availability'])),
              'missing_source_rows': {p: int(joined[p + '_source'].isna().sum()) for p in frames},
              'duplicate_keys': int(joined.duplicated(keys).sum()),
              'join': 'one-to-one full outer join on official date + unchanged hospital code',
              'limitations': '同日匹配不代表同一批患者或同一时间窗口；BOR为官方所属日期结束后的午夜清点，详细床位分母待核实。'}
WORK.mkdir(parents=True, exist_ok=True)
payload = {'datasets': datasets, 'joined': {'fields': join_fields, 'rows': json_rows(joined, join_fields), 'audit': join_audit}}
(WORK / 'hospital_metrics.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'sources': audits, 'joined': join_audit}, ensure_ascii=True, indent=2))
