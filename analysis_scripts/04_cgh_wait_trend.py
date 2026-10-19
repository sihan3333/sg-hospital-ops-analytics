"""CGH 等待趋势教学图：完整自然日历、每日 P50、严格完整的 7 日移动均值。"""
from datetime import datetime
from hashlib import sha256
import importlib.metadata
import json
import os
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / '处理脚本'
WORK = SCRIPTS / '.work_trend'
OUT = ROOT / 'outputs/01a0bcf0-303c-7d21-9f82-aafceef76c21'
WORK.mkdir(parents=True, exist_ok=True)
OUT.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(SCRIPTS / '.plot_packages'))
os.environ['MPLCONFIGDIR'] = str(WORK / 'matplotlib_cache')

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from matplotlib.font_manager import FontProperties, fontManager
from openpyxl import load_workbook

source = ROOT / '原始数据/2026-09-20/入院等待时间_WT_week36Y2026.xlsx'
payload = json.loads((SCRIPTS / '.work_join/hospital_metrics.json').read_text(encoding='utf-8'))
assert sha256(source.read_bytes()).hexdigest() == payload['datasets']['wait']['audit']['sha256']
frame = pd.DataFrame(payload['joined']['rows'], columns=payload['joined']['fields'])
cgh = frame.loc[frame.hospital_code.eq('CGH')].copy()
cgh['date'] = pd.to_datetime(cgh['date'])
assert not cgh['date'].duplicated().any()
daily = cgh.set_index('date')['admission_wait_p50_hours'].sort_index().asfreq('D')
weekly = daily.rolling(window=7, min_periods=7).mean()

# 独立重读原表：不能只用派生表验证派生表。
book = load_workbook(source, read_only=True, data_only=True)
original = {}
for name in ('Historical', 'Sheet1'):
    sheet = book[name]
    header = next(sheet.iter_rows(min_row=3, max_row=3, max_col=10, values_only=True))
    cgh_col = header.index('CGH')
    for row in sheet.iter_rows(min_row=4, max_col=10, values_only=True):
        if isinstance(row[0], datetime) and row[0] >= datetime(2024, 1, 1):
            key = pd.Timestamp(row[0])
            assert key not in original
            original[key] = row[cgh_col]
book.close()
for day, value in daily.items():
    if day not in original or original[day] is None:
        assert pd.isna(value)
    else:
        assert value == original[day]
assert set(daily.dropna().index) == set(original)

# 逐个窗口核查：缺一天即不可计算，避免将缺失前后的“7条记录”当成7个自然日。
for i, day in enumerate(daily.index):
    window = daily.iloc[max(0, i - 6):i + 1]
    if len(window) != 7 or window.isna().any():
        assert pd.isna(weekly.loc[day])
    else:
        assert np.isclose(weekly.loc[day], sum(float(v) for v in window) / 7, rtol=1e-12)
assert daily.loc['2026-08-02':'2026-08-08'].isna().all()
assert weekly.loc['2026-08-02':'2026-08-14'].isna().all()
assert pd.notna(weekly.loc['2026-08-15'])

peak_day = daily.idxmax()
week_end = weekly.idxmax()
week_start = week_end - pd.Timedelta(days=6)
prior_start = week_start - pd.Timedelta(days=7)
prior_end = week_start - pd.Timedelta(days=1)
prior_mean = float(daily.loc[prior_start:prior_end].mean())
spike_day = pd.Timestamp('2024-09-14')
audit = {
    'hospital': 'CGH', 'calendar_days': len(daily), 'observed_days': int(daily.notna().sum()),
    'valid_seven_day_windows': int(weekly.notna().sum()),
    'missing_dates': [d.date().isoformat() for d in daily.index[daily.isna()]],
    'max_daily': {'date': peak_day.date().isoformat(), 'hours': float(daily.max())},
    'max_rolling_window': {'start': week_start.date().isoformat(), 'end': week_end.date().isoformat(),
                           'mean_of_daily_p50_hours': float(weekly.max()),
                           'daily_p50_hours': [float(v) for v in daily.loc[week_start:week_end]]},
    'previous_window': {'start': prior_start.date().isoformat(), 'end': prior_end.date().isoformat(), 'mean_of_daily_p50_hours': prior_mean},
    'local_spike': {'date': spike_day.date().isoformat(), 'hours': float(daily.loc[spike_day]),
                    'previous_day_hours': float(daily.loc[spike_day-pd.Timedelta(days=1)]),
                    'next_day_hours': float(daily.loc[spike_day+pd.Timedelta(days=1)])},
    'definition': '当日及前6个自然日的每日P50等权平均，7天全部非空才计算；不是患者总体均值或总体中位数。',
    'selection': '两段局部图为事后选择的读图例子，不是预先设定的异常检测规则或因果证据。',
    'source_sha256': sha256(source.read_bytes()).hexdigest(),
}
(WORK / 'cgh_trend_audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')

font_file = Path('C:/Windows/Fonts/msyh.ttc')
fontManager.addfont(str(font_file))
font_name = FontProperties(fname=str(font_file)).get_name()
plt.rcParams.update({'font.family': font_name, 'axes.unicode_minus': False, 'font.size': 11,
                     'axes.spines.top': False, 'axes.spines.right': False, 'svg.fonttype': 'path',
                     'savefig.facecolor': '#FFFFFF'})
fig = plt.figure(figsize=(15, 10), facecolor='white')
grid = fig.add_gridspec(2, 2, height_ratios=[1.15, 1], left=.075, right=.97, top=.855, bottom=.19, hspace=.53, wspace=.17)
overview = fig.add_subplot(grid[0, :])
left = fig.add_subplot(grid[1, 0])
right = fig.add_subplot(grid[1, 1])
daily_color, smooth_color = '#6799BF', '#C65B20'


def plot_series(ax, start, end, markers=False):
    actual = daily.loc[start:end]
    trend = weekly.loc[start:end]
    ax.plot(actual.index, actual, color=daily_color, lw=1.1 if not markers else 1.8,
            alpha=.65 if not markers else 1, marker='o' if markers else None, ms=4,
            label='每日入院等待中位数（P50）')
    ax.plot(trend.index, trend, color=smooth_color, lw=2.2, label='过去7天每日P50的平均值')
    ax.set_ylim(0, 32)
    ax.set_xlim(pd.Timestamp(start), pd.Timestamp(end))
    ax.set_yticks([0, 5, 10, 15, 20, 25, 30])
    ax.set_ylabel('小时', rotation=0, labelpad=18)
    ax.grid(axis='y', color='#E4E9EF', lw=.8)
    ax.set_axisbelow(True)
    ax.spines['left'].set_color('#CBD5E1')
    ax.spines['bottom'].set_color('#CBD5E1')
    ax.tick_params(axis='both', colors='#435366', length=3)


plot_series(overview, daily.index.min(), daily.index.max())
overview.set_title('全期走势', loc='left', fontsize=12, fontweight='bold', pad=12)
overview.xaxis.set_major_locator(mdates.MonthLocator(bymonth=[1, 5, 9]))
overview.xaxis.set_major_formatter(mdates.DateFormatter('%Y-%m'))
overview.legend(loc='upper right', frameon=False, ncol=2, fontsize=10)
overview.axvspan(pd.Timestamp('2026-08-02'), pd.Timestamp('2026-08-09'), color='#D9DEE5', alpha=.9, zorder=0)
overview.annotate('缺失7天', xy=(pd.Timestamp('2026-08-05'), 10), xytext=(pd.Timestamp('2026-04-15'), 20),
                  fontsize=10, color='#5B6573', arrowprops={'arrowstyle':'->', 'color':'#8B95A3', 'lw':1})

plot_series(left, '2024-04-20', '2024-05-20', True)
left.axvspan(week_start, week_end+pd.Timedelta(days=1), color='#FBEADB', alpha=.7, zorder=0)
left.set_title('A  连续多日较高：2024年5月', loc='left', fontsize=12, fontweight='bold', pad=14)
left.xaxis.set_major_locator(mdates.DayLocator(interval=7))
left.xaxis.set_major_formatter(mdates.DateFormatter('%m-%d'))
left.annotate('5月2日：29.00小时', xy=(peak_day, daily.loc[peak_day]), xytext=(pd.Timestamp('2024-04-21'), 31),
              fontsize=10, color='#274B71', arrowprops={'arrowstyle':'->', 'color':'#6799BF', 'lw':1})
left.text(.98, .08, f'5/2—5/8：7天均值 {weekly.max():.2f} 小时\n此前7天：{prior_mean:.2f} 小时',
          transform=left.transAxes, ha='right', va='bottom', fontsize=10, color='#8D431B',
          bbox={'boxstyle':'round,pad=.4', 'facecolor':'white', 'edgecolor':'#E6D8CB', 'alpha':.94})

plot_series(right, '2024-09-07', '2024-09-21', True)
right.set_title('B  单日尖峰：2024年9月', loc='left', fontsize=12, fontweight='bold', pad=14)
right.xaxis.set_major_locator(mdates.DayLocator(interval=3))
right.xaxis.set_major_formatter(mdates.DateFormatter('%m-%d'))
right.annotate('9月14日：22.75小时', xy=(spike_day, daily.loc[spike_day]), xytext=(pd.Timestamp('2024-09-08'), 29),
               fontsize=10, color='#274B71', arrowprops={'arrowstyle':'->','color':'#6799BF','lw':1})
right.annotate('次日：5.42小时', xy=(spike_day+pd.Timedelta(days=1), daily.loc[spike_day+pd.Timedelta(days=1)]),
               xytext=(pd.Timestamp('2024-09-16'), 14), fontsize=10, color='#274B71',
               arrowprops={'arrowstyle':'->','color':'#6799BF','lw':1})

fig.text(.075, .963, 'CGH 入院等待时间：区分单日波动与持续变化', fontsize=19, fontweight='bold', color='#203D5B')
fig.text(.075, .925, '2024-01-01 至 2026-09-12  |  979个有记录日期，7天缺失  |  同一家医院的描述性趋势',
         fontsize=11, color='#52657A')
fig.text(.075, .12, '橙线：当日及此前6个自然日的每日P50等权平均；不代表七天全部患者的平均等待或中位数。', fontsize=10, color='#435366')
fig.text(.075, .09, '7天齐全才计算。原始记录缺2026-08-02至08-08；橙线在8月15日恢复。局部例子不构成告警阈值或原因判断。', fontsize=10, color='#435366')
fig.text(.075, .055, '来源：MOH，WT for Admission to Ward_week36Y2026.xlsx；下载于2026-09-20。', fontsize=9.5, color='#65758A')
fig.savefig(OUT / 'CGH_入院等待趋势.png', dpi=170)
plt.close(fig)

window_rows = '\n'.join(f'| {d:%Y-%m-%d} | {v:.2f} |' for d, v in daily.loc[week_start:week_end].items())
report = f'''# CGH 入院等待趋势：第一张图

研究期：2024-01-01 至 2026-09-12。数据粒度：一家医院一天的入院等待时间中位数，单位为小时。

## 方法

从已核查的合并底表筛选 CGH，按日期排序并保留完整日历，共 {len(daily)} 天，其中 {daily.notna().sum()} 天有等待记录。蓝线直接显示每日 P50；橙线为当日及此前 6 个自然日的每日 P50 等权平均。选择 7 天是为了用一个完整星期的窗口观察较平滑的走势，不是预测模型，也不构成异常阈值。

连续 7 天全部有值才计算均值，共 {weekly.notna().sum()} 个可用窗口。开头 6 天不计算；2026-08-02 至 08-08 的原始缺失使 7 天线在 08-02 至 08-14 不可用，08-15 恢复。绘图保留 NaN，让折线断开，不填 0、不跨缺口连线。

注意：橙线是每日中位数的平均，不是七天所有患者的平均等待时间或中位数。一个单日尖峰也会影响随后多个重叠窗口，必须与蓝线一起看。

## 读图例子

两段例子为事后选取的教学窗口，不是经过验证的告警规则。

A：{week_start:%Y-%m-%d} 至 {week_end:%Y-%m-%d}，每日 P50 的七日均值为 {weekly.max():.2f} 小时，此前 {prior_start:%m-%d} 至 {prior_end:%m-%d} 为 {prior_mean:.2f} 小时。全期最高单日为 {peak_day:%Y-%m-%d}，P50 为 {daily.max():.2f} 小时；同一周其他日期也多次处于较高水平，因此不只看最高的一天。

| 日期 | 每日 P50（小时） |
|---|---:|
{window_rows}

B：2024-09-14 的 P50 为 22.75 小时；前一天约 9.77 小时，后一天为 5.42 小时，再后一天约 2.77 小时。这段适合观察局部单日尖峰和随后回落，不能据此排除该日真实的服务压力。

本课仅识别值得继续核查的时间段；尚未结合急诊量、BOR、假期或院内事件判断原因。尚未衡量业务策略采纳、执行或等待改善效果。

## 验证与来源

每日数值已独立重读原始等待 Excel 比对；每个移动窗口均按自然日重新核算，检查了缺失周断点及恢复日期。原件 SHA256 未改变。

来源：[MOH 入院等待统计页面](https://www.moh.gov.sg/others/resources-and-statistics/healthcare-institution-statistics-waiting-time-for-admission-to-ward/)。原件为本项目 2026-09-20 下载的 WT for Admission to Ward_week36Y2026.xlsx。
'''
(OUT / 'CGH_入院等待趋势_读图说明.md').write_text(report, encoding='utf-8')
versions = {name: importlib.metadata.version(name) for name in ['matplotlib','numpy','pandas','pyparsing','contourpy','cycler','fonttools','kiwisolver']}
(WORK / 'plot_versions.json').write_text(json.dumps(versions, indent=2), encoding='utf-8')
print(json.dumps(audit, ensure_ascii=True, indent=2))
print(json.dumps({'png': str(OUT / 'CGH_入院等待趋势.png'), 'method_notes': str(OUT / 'CGH_入院等待趋势_读图说明.md')}, ensure_ascii=True))
