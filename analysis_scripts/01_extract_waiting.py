"""第一步：只读原表，拼接日期行，检查覆盖范围；不填补或改写医院数值。"""
from collections import Counter
from datetime import date, datetime, timedelta
from hashlib import sha256
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / '原始数据/2026-09-20/入院等待时间_WT_week36Y2026.xlsx'
WORK = ROOT / '处理脚本/.work_waiting'
START = date(2024, 1, 1)

before = sha256(SOURCE.read_bytes()).hexdigest()
book = load_workbook(SOURCE, read_only=True, data_only=True)
records, sheet_stats = [], []
expected_headers = ['Date', 'AH', 'CGH', 'KTPH', 'NTFGH', 'NUH(A)', 'SGH', 'SKH', 'TTSH', 'WH']
for name in ('Historical', 'Sheet1'):
    sheet = book[name]
    headers = next(sheet.iter_rows(min_row=3, max_row=3, max_col=10, values_only=True))
    if list(headers) != expected_headers:
        raise ValueError(f'{name}: 表头改变，需重新检查')
    dates = []
    for source_row, values in enumerate(sheet.iter_rows(min_row=4, max_col=10, values_only=True), 4):
        raw_date = values[0]
        # 排除页尾来源说明和纯空白行，只取带真实日期的行。
        if not isinstance(raw_date, (date, datetime)):
            continue
        day = raw_date.date() if isinstance(raw_date, datetime) else raw_date
        dates.append(day)
        if day < START:
            continue
        if any(v is not None and not isinstance(v, (int, float)) for v in values[1:]):
            raise ValueError(f'{name}!{source_row}: 存在非数值，停止自动整理')
        records.append([day.isoformat(), *values[1:], name, source_row])
    sheet_stats.append({'sheet': name, 'all_date_rows': len(dates),
                        'start': min(dates).isoformat(), 'end': max(dates).isoformat(),
                        'study_date_rows': sum(d >= START for d in dates)})
book.close()
records.sort(key=lambda row: row[0])
counts = Counter(row[0] for row in records)
duplicates = {day: n for day, n in counts.items() if n > 1}
if duplicates:
    # 不擅自保留第一条、取均值或删除冲突记录。
    raise ValueError(f'发现重复日期，须先检查：{duplicates}')
end = date.fromisoformat(records[-1][0])
calendar = [START + timedelta(days=i) for i in range((end - START).days + 1)]
missing = [day.isoformat() for day in calendar if day.isoformat() not in counts]
blank_counts = {h: sum(row[i] is None for row in records) for i, h in enumerate(expected_headers[1:], 1)}
zero_counts = {h: sum(row[i] == 0 for row in records) for i, h in enumerate(expected_headers[1:], 1)}
audit = {'study_start': START.isoformat(), 'study_end': end.isoformat(),
         'source_sha256': before, 'sheet_stats': sheet_stats,
         'actual_date_rows': len(records), 'expected_date_rows': len(calendar),
         'duplicate_dates': duplicates, 'missing_dates': missing,
         'blank_cells_by_hospital': blank_counts, 'zero_cells_by_hospital': zero_counts,
         'policy': '保留原值、空白及数值0；不补行、不插值、不求均值；日期覆盖检查仅针对本次源文件。'}
if sha256(SOURCE.read_bytes()).hexdigest() != before:
    raise RuntimeError('运行期间源文件发生改变')
WORK.mkdir(parents=True, exist_ok=True)
(WORK / 'waiting_data.json').write_text(json.dumps({'audit': audit, 'headers': expected_headers, 'rows': records}, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(audit, ensure_ascii=True, indent=2))
